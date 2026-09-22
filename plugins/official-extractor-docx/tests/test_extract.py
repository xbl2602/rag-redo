from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

_PLUGIN_DIR = Path(__file__).parent.parent
_REPO_ROOT = _PLUGIN_DIR.parent.parent
for p in (_REPO_ROOT, _PLUGIN_DIR):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

import docx  # noqa: E402

from official_extractor_docx.extract import extract  # noqa: E402


def _make_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("标题一", level=1)
    document.add_paragraph("第一段正文内容。")
    document.add_heading("子标题", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列A"
    table.cell(0, 1).text = "列B"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    document.add_paragraph("表格后面的段落。")
    document.save(path)


class TestExtractDocx(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def test_extracts_headings_paragraphs_and_table(self):
        _make_docx(self.tmp / "a.docx")
        doc = extract("lib1", "a.docx", self.tmp)
        self.assertIsNone(doc.failure_reason)
        self.assertIn("# 标题一", doc.text)
        self.assertIn("## 子标题", doc.text)
        self.assertIn("第一段正文内容。", doc.text)
        self.assertIn("| 列A | 列B |", doc.text)
        self.assertIn("值1", doc.text)

    def test_order_preserved_paragraph_table_paragraph(self):
        _make_docx(self.tmp / "a.docx")
        doc = extract("lib1", "a.docx", self.tmp)
        before_table = doc.text.index("子标题")
        table_pos = doc.text.index("列A")
        after_pos = doc.text.index("表格后面的段落")
        self.assertLess(before_table, table_pos)
        self.assertLess(table_pos, after_pos)

    def test_empty_document_folds_to_failure(self):
        docx.Document().save(self.tmp / "empty.docx")
        doc = extract("lib1", "empty.docx", self.tmp)
        self.assertIsNone(doc.text)
        self.assertIn("空", doc.failure_reason)

    def test_missing_file_folds_to_failure(self):
        doc = extract("lib1", "missing.docx", self.tmp)
        self.assertIsNone(doc.text)
        self.assertIsNotNone(doc.failure_reason)

    def test_corrupted_docx_folds_to_failure_not_exception(self):
        (self.tmp / "bad.docx").write_bytes(b"not a real docx/zip")
        doc = extract("lib1", "bad.docx", self.tmp)  # 不应该抛异常
        self.assertIsNone(doc.text)
        self.assertIsNotNone(doc.failure_reason)


if __name__ == "__main__":
    unittest.main()
