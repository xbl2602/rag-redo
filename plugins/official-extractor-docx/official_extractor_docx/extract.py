"""DOCX 的提取逻辑：读文件 -> ExtractedDocument（结构化转 Markdown）。

标题样式（Heading 1/2/3...）转成对应级别的 `#`，普通段落原样输出，表格
转成 Markdown 管道表格——这样切块阶段（official-chunker）能复用同一套
"按标题分段、表格不拆开"的逻辑，不用为 DOCX 来源的文档另写一套处理。
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import docx

from core.contracts import ExtractedDocument

EXTRACTOR_VERSION = "0.1.0"
PLUGIN_ID = "official-extractor-docx"

_HEADING_LEVEL_RE = re.compile(r"^Heading (\d+)$")


def _content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _fail(library_id: str, path: str, reason: str, content_hash: str = "") -> ExtractedDocument:
    return ExtractedDocument(
        library_id=library_id,
        path=path,
        text=None,
        failure_reason=reason,
        extracted_by=PLUGIN_ID,
        extractor_version=EXTRACTOR_VERSION,
        content_hash=content_hash,
    )


def _table_to_markdown(table: "docx.table.Table") -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _document_to_markdown(document: "docx.Document") -> str:
    """按 body 的实际元素顺序遍历段落和表格，不能简单先取全部段落再取
    全部表格——那样会打乱"表格出现在哪两段文字之间"的原始顺序。"""
    from docx.oxml.ns import qn  # noqa: PLC0415 - 局部import，避免给模块顶层加不必要的内部符号依赖

    parts: list[str] = []
    body = document.element.body
    paragraphs_by_id = {p._p: p for p in document.paragraphs}
    tables_by_id = {t._tbl: t for t in document.tables}

    for child in body.iterchildren():
        if child.tag == qn("w:p") and child in paragraphs_by_id:
            para = paragraphs_by_id[child]
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style is not None else ""
            m = _HEADING_LEVEL_RE.match(style_name or "")
            if m:
                level = min(int(m.group(1)), 6)
                parts.append(f"{'#' * level} {text}")
            else:
                parts.append(text)
        elif child.tag == qn("w:tbl") and child in tables_by_id:
            md_table = _table_to_markdown(tables_by_id[child])
            if md_table:
                parts.append(md_table)

    return "\n\n".join(parts)


def extract(library_id: str, path: str, root: Path) -> ExtractedDocument:
    full_path = root / path
    try:
        data = full_path.read_bytes()
    except OSError as exc:
        return _fail(library_id, path, f"读取失败: {type(exc).__name__}: {exc}")

    content_hash = _content_hash(data)

    try:
        document = docx.Document(str(full_path))
        text = _document_to_markdown(document)
    except Exception as exc:  # noqa: BLE001 - extractor 绝不抛异常
        return _fail(library_id, path, f"提取失败: {type(exc).__name__}: {exc}", content_hash)

    if not text.strip():
        return _fail(library_id, path, "提取结果为空", content_hash)

    return ExtractedDocument(
        library_id=library_id,
        path=path,
        text=text,
        failure_reason=None,
        extracted_by=PLUGIN_ID,
        extractor_version=EXTRACTOR_VERSION,
        content_hash=content_hash,
    )
